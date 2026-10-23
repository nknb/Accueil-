"""Données corrompues ou extrêmes: aucune page ne doit planter (pytest-qt échoue si un slot lève une exception)."""
import uuid
from datetime import date

import pytest

pytest.importorskip("PyQt6.QtCharts")

from database.connection import db  # noqa: E402

AUJOURD_HUI = date.today().isoformat()
LONG = "Ñ" * 5000


def _inserer(**champs):
    base = {"id": uuid.uuid4().hex, "date": AUJOURD_HUI, "nom": "N", "prenom": "P", "statut": "en_attente",
            "created_at": f"{AUJOURD_HUI}T10:00:00", "updated_at": f"{AUJOURD_HUI}T10:00:00"}
    base.update(champs)
    db.execute_update(f"INSERT INTO accueil_demandeaccueil ({', '.join(base)}) VALUES ({', '.join('?' * len(base))})", tuple(base.values()))


@pytest.fixture
def base_hostile():
    db.execute_update("DELETE FROM accueil_demandeaccueil")
    _inserer(nom=None, prenom=None)                                                # champs vides
    _inserer(nom=LONG, prenom=LONG, numero_cnps=LONG, contacts=LONG)               # textes énormes
    _inserer(nom="😀 ÉLODIE <b>gras</b> '; DROP TABLE x;--", prenom="\x00nul", contacts="=1+1")   # emoji, HTML, SQL, formule
    _inserer(heure_arrivee="pas une heure", heure_reception="zz:zz", temps_attente="x:y:z", satisfaction="beaucoup")
    _inserer(heure_arrivee="2026-13-45T99:99:99", temps_attente="12", satisfaction=99, delai_objectif="oui")
    _inserer(statut="statut_inconnu", priorite="urgentissime", temps_attente="9999:99:99")
    _inserer(statut="traite", heure_arrivee="23:59:59", heure_reception="00:00:01", temps_attente="0:00:02", satisfaction=3)   # réception après minuit
    _inserer(date="pas une date", statut="traite")                                 # date corrompue
    _inserer(date=AUJOURD_HUI + "T08:30:00", statut="non_traite", agent_accueil_id=424242)   # agent inexistant
    return 9


def test_pages_accueil_et_preaccueil_ne_plantent_pas(qtbot, base_hostile):
    from widgets.accueil import AccueilPage
    from widgets.preaccueil import PreAccueilPage
    accueil, pre = AccueilPage(), PreAccueilPage()
    for p in (accueil, pre):
        qtbot.addWidget(p)
    accueil.set_current_user({"id": 1, "role": "accueil"})
    qtbot.waitUntil(lambda: accueil.ui.tableWidget.rowCount() >= 1, timeout=10000)
    qtbot.waitUntil(lambda: pre.ui.tableWidget.rowCount() >= 1, timeout=10000)
    for ligne in range(accueil.ui.tableWidget.rowCount()):                         # sélection de chaque ligne
        accueil.ui.tableWidget.setCurrentCell(ligne, 0)
    accueil.grab()
    pre.grab()


def test_page_base_de_donnees_ne_plante_pas(qtbot, base_hostile):
    from widgets.database import DatabasePage
    page = DatabasePage()
    qtbot.addWidget(page)
    qtbot.waitUntil(lambda: page.table.rowCount() >= 1, timeout=10000)
    page.grab()
    for terme in ("élodie", "'; DROP", "%", "_", "\\", "😀", ""):
        page._global_search = terme
        page._apply_filters(reset_page=True)
        qtbot.wait(50)
    assert db.execute_query("SELECT COUNT(*) AS n FROM accueil_demandeaccueil")[0]["n"] == base_hostile   # rien n'a été supprimé


def test_dashboard_ne_plante_pas_et_compte_juste(qtbot, base_hostile):
    from widgets.dashboard.Main import DashboardPage, _DashboardWorker
    page = DashboardPage()
    qtbot.addWidget(page)
    qtbot.waitUntil(lambda: page._chart_handler._last_data is not None, timeout=15000)
    page.grab()
    res = _DashboardWorker("2000-01-01", "2100-12-31").load()
    assert res["total_demandes"] >= base_hostile - 1                              # la date corrompue peut être exclue de la période


def test_export_excel_avec_donnees_hostiles(qtbot, base_hostile, tmp_path):
    import openpyxl
    from widgets.dashboard.Main import _ChartHandler, _DashboardWorker
    data = _DashboardWorker("2000-01-01", "2100-12-31").load()
    sortie = tmp_path / "r.xlsx"
    _ChartHandler._write_excel(None, str(sortie), data, "2000-01-01", "2100-12-31")
    brut = openpyxl.load_workbook(sortie)["Données brutes"]
    assert all(c.data_type != "f" for ligne in brut.iter_rows() for c in ligne)   # aucune formule injectée


def test_sauvegarde_et_jeton_sur_base_hostile(base_hostile, tmp_path):
    from database.sauvegarde import sauvegarder_base
    assert db.change_token()[0] >= 0
    copie = sauvegarder_base(db.db_path, dossier=tmp_path)
    assert copie.stat().st_size > 0


def test_nettoyer_texte_retire_les_caracteres_de_controle_seulement():
    from utils.validators import nettoyer_texte
    assert nettoyer_texte("A\x00B\x0bC\x1fD") == "ABCD"
    assert nettoyer_texte("Élodie N'Guessan\tÀ\nB\r") == "Élodie N'Guessan\tÀ\nB\r"      # accents, tabulation et sauts de ligne conservés
    assert nettoyer_texte(None) == "" and nettoyer_texte(12) == "12"


def test_preaccueil_ne_stocke_pas_de_caracteres_de_controle(qtbot, monkeypatch):
    from PyQt6.QtWidgets import QMessageBox
    from widgets.preaccueil import PreAccueilPage
    for nom in ("information", "warning", "critical"):
        monkeypatch.setattr(QMessageBox, nom, staticmethod(lambda *a, **k: None))
    db.execute_update("DELETE FROM accueil_demandeaccueil")
    page = PreAccueilPage()
    qtbot.addWidget(page)
    page.input_nom.setText("Ko\x0bné")           # collé depuis un document: tabulation verticale invisible
    page.input_prenom.setText("Ami\x00nata")
    page.input_contact.setText("07\x1f00")
    page.combo_service.setCurrentIndex(1)
    page.combo_motif.setCurrentIndex(1)
    page.enregistrer_demande()
    ligne = db.execute_query("SELECT nom, prenom, contacts FROM accueil_demandeaccueil")[0]
    assert (ligne["nom"], ligne["prenom"], ligne["contacts"]) == ("KONÉ", "Aminata", "0700")


def test_export_de_la_page_base_de_donnees_survit_aux_caracteres_de_controle(qtbot, base_hostile, tmp_path, monkeypatch):
    import openpyxl
    from PyQt6.QtWidgets import QFileDialog, QMessageBox
    from widgets.database import DatabasePage
    sortie = tmp_path / "bd.xlsx"
    messages = []
    monkeypatch.setattr(QFileDialog, "getSaveFileName", staticmethod(lambda *a, **k: (str(sortie), "")))
    for nom in ("information", "critical", "warning"):
        monkeypatch.setattr(QMessageBox, nom, staticmethod(lambda *a, _n=nom, **k: messages.append((_n, a[2]))))
    page = DatabasePage()
    qtbot.addWidget(page)
    qtbot.waitUntil(lambda: page.table.rowCount() >= 1, timeout=10000)
    page._on_export_excel()
    assert [m for m in messages if m[0] == "critical"] == []          # pas d'« Erreur lors de l'export »
    assert sortie.exists()
    feuille = openpyxl.load_workbook(sortie).active
    assert feuille.max_row >= base_hostile + 1                          # en-tête + toutes les demandes


@pytest.mark.parametrize("identifiant,mot_de_passe", [
    ("' OR '1'='1' --", "x"),
    ("admin'; DROP TABLE auth_user;--", "x"),
    ("a" * 500, "b" * 500),                       # bien plus que les 72 octets que bcrypt accepte
    ("é😀\x00", "mot\x00de\x1fpasse"),
])
def test_connexion_refuse_les_identifiants_hostiles_sans_planter(qtbot, monkeypatch, identifiant, mot_de_passe):
    from PyQt6.QtWidgets import QMessageBox
    import main
    monkeypatch.setattr(QMessageBox, "information", staticmethod(lambda *a, **k: None))
    avant = db.execute_query("SELECT COUNT(*) AS n FROM auth_user")[0]["n"]
    fenetre = main.LoginApp()
    qtbot.addWidget(fenetre)
    ouvertures = []
    monkeypatch.setattr(fenetre, "ouvrir_fenetre_principale", lambda: ouvertures.append(1))
    fenetre.verifier_db(identifiant, mot_de_passe)
    assert not ouvertures and fenetre.utilisateur is None
    assert db.execute_query("SELECT COUNT(*) AS n FROM auth_user")[0]["n"] >= avant      # la table n'a pas été touchée


def test_fenetre_de_traitement_s_ouvre_sur_toute_demande(qtbot, base_hostile):
    from widgets.accueil.accueil_page import TraitementDialog
    lignes = db.execute_query(
        "SELECT d.id, d.nom, d.prenom, d.numero_cnps, d.contacts, d.heure_arrivee, d.satisfaction, d.commentaire_satisfaction,"
        " d.statut, s.nom AS service, m.nom AS motif FROM accueil_demandeaccueil d"
        " LEFT JOIN accueil_service s ON d.service_demande_id = s.id LEFT JOIN accueil_motif m ON d.motif_id = m.id")
    assert len(lignes) == base_hostile
    for ligne in lignes:
        dialogue = TraitementDialog(ligne)
        qtbot.addWidget(dialogue)
        p = dialogue.payload()
        assert set(p) == {"heure_reception", "satisfaction", "commentaire_satisfaction"}
        dialogue.grab()


def test_fiche_usager_du_dashboard_sur_donnees_hostiles(qtbot, base_hostile):
    from widgets.dashboard.Main import _DashboardWorker, _UserDetailDialog
    usagers = _DashboardWorker("2000-01-01", "2100-12-31").load()["top_users"]
    assert usagers
    for usager in usagers:
        fiche = _UserDetailDialog(usager)
        qtbot.addWidget(fiche)
        fiche.grab()


# ── Partage réseau indisponible ───────────────────────────────────────────────

def test_base_inaccessible_ne_plante_pas_au_demarrage(tmp_path):
    from database import connection
    absent = tmp_path / "partage_absent" / "dossier" / "db.sqlite3"           # dossier qui n'existe pas
    d = connection.DatabaseConnection(str(absent))                            # ne doit pas lever
    assert d.change_token() == connection._ERROR_TOKEN
    with pytest.raises(Exception):
        d.execute_query("SELECT 1")                                           # l'erreur est explicite, pas un blocage
    assert not absent.exists()                                                # aucune base fantôme créée au mauvais endroit


def test_chemin_unc_inaccessible_donne_un_message_clair():
    from database import connection
    d = connection.DatabaseConnection("//serveur-qui-n-existe-pas/partage/accueil/db.sqlite3")
    with pytest.raises(RuntimeError, match="Partage réseau inaccessible"):
        d.execute_query("SELECT 1")
    assert d.change_token() == connection._ERROR_TOKEN


def test_connexion_avec_base_inaccessible_affiche_le_message_reseau(qtbot, monkeypatch, tmp_path):
    import main
    from PyQt6.QtWidgets import QMessageBox
    from config import AppConfig
    monkeypatch.setattr(QMessageBox, "information", staticmethod(lambda *a, **k: None))
    monkeypatch.setattr(AppConfig, "get_db_path", classmethod(lambda c: tmp_path / "absent" / "db.sqlite3"))
    fenetre = main.LoginApp()                                                 # ne doit pas planter à l'ouverture
    qtbot.addWidget(fenetre)
    fenetre.verifier_db("admin", "x")
    assert fenetre.utilisateur is None
    assert "Réseau indisponible" in fenetre.message_label.text()


def test_dossier_de_base_absent_echoue_vite_sans_reessayer(tmp_path):
    """« unable to open » est réessayé (glitch réseau passager), mais pas si le dossier n'existe pas: réessayer n'y change rien."""
    import time
    from database import connection
    d = connection.DatabaseConnection(str(tmp_path / "absent" / "db.sqlite3"))
    t0 = time.perf_counter()
    for _ in range(3):
        with pytest.raises(Exception):
            d.execute_query("SELECT 1")
    assert time.perf_counter() - t0 < 1.0
