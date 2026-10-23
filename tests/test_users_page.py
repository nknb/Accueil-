"""Garde-fous de la page Utilisateurs: dernier admin."""
from database.connection import db


def _add(username, role):
    is_admin = 1 if role == "admin" else 0
    return db.execute_insert(
        "INSERT INTO auth_user (username, password, role, is_active, is_superuser, is_staff, date_joined)"
        " VALUES (?, 'x', ?, 1, ?, ?, '2000-01-01')", (username, role, is_admin, is_admin))


def test_dernier_admin_est_protege():
    from widgets.users.Main import retire_admin_autorise
    db.execute_update("DELETE FROM auth_user")
    a1 = _add("adm1", "admin")
    agent = _add("agent1", "accueil")
    assert retire_admin_autorise(agent)          # pas un admin: libre
    assert not retire_admin_autorise(a1)         # seul admin actif: interdit
    a2 = _add("adm2", "admin")
    assert retire_admin_autorise(a1)             # un autre admin existe
    assert retire_admin_autorise(a2)


def test_dashboard_compte_par_date_de_la_demande_pas_par_modification(qtbot):
    from datetime import date
    from database.models import DemandeAccueil
    from widgets.dashboard.Main import _DashboardWorker
    jour = "2011-05-06"
    DemandeAccueil(nom="DASH", date=date(2011, 5, 6)).save()
    db.execute_update("DELETE FROM accueil_demandeaccueil WHERE date != ?", (jour,))
    db.rollover_pending_to_non_traite()  # réécrit updated_at à maintenant
    res = []
    w = _DashboardWorker(jour, jour)
    w.finished.connect(res.append)
    w.run()
    assert res and res[0]["total_demandes"] == 1


def test_dashboard_temps_attente_format_h_mm_ss(qtbot):
    """L'appli écrit 'H:MM:SS' (ex. 0:12:30), pas 'HH:MM:SS': 12 min ne doit pas devenir 2 min."""
    from datetime import date
    from database.models import DemandeAccueil
    from widgets.dashboard.Main import _DashboardWorker
    jour = date(2012, 6, 7)
    db.execute_update("DELETE FROM accueil_demandeaccueil")
    for t in ("0:12:30", "0:45:10", "00:20:00"):
        DemandeAccueil(nom="W", date=jour, temps_attente=t).save()
    res = []
    w = _DashboardWorker("2012-06-07", "2012-06-07")
    w.finished.connect(res.append)
    w.run()
    d = dict(zip(res[0]["wait_labels"], res[0]["wait_counts"]))
    assert d["≤ 15 min"] == 1 and d["15–30 min"] == 1 and d["30–60 min"] == 1
    assert round(res[0]["avg_wait_min"], 1) == 25.7  # (12 + 45 + 20) / 3


def test_helpers_dashboard_attente_et_heure():
    from widgets.dashboard.Main import _minutes_attente, _heure_arrivee
    assert _minutes_attente("0:12:30") == 12 and _minutes_attente("01:05:00") == 65
    assert _minutes_attente("") is None and _minutes_attente("abc") is None
    assert _heure_arrivee("2026-10-23T10:50:03.483938") == 10
    assert _heure_arrivee("09:15:00") == 9 and _heure_arrivee(None) is None


def test_export_excel_pas_de_formule_et_totaux_complets(tmp_path):
    import openpyxl
    from datetime import date
    from database.models import DemandeAccueil
    from widgets.dashboard.Main import _ChartHandler
    db.execute_update("DELETE FROM accueil_demandeaccueil")
    jour = date(2013, 7, 8)
    DemandeAccueil(nom="=1+1", prenom="@x", numero_cnps="=SUM(1,2)", date=jour).save()
    data = {"total_demandes": 12, "motifs_names": ["A"], "motifs_counts": [5],
            "services_names": ["S"], "services_counts": [5]}  # 7 demandes hors top
    out = tmp_path / "r.xlsx"
    _ChartHandler._write_excel(None, str(out), data, "2013-07-08", "2013-07-08")
    wb = openpyxl.load_workbook(out)
    brut = wb["Données brutes"]
    assert all(c.data_type != "f" for row in brut.iter_rows() for c in row)
    assert brut["B3"].value == "=1+1"            # texte conservé, pas évalué
    mot = [[c.value for c in r] for r in wb["Motifs"].iter_rows(min_row=4)]
    assert ["Autres", 7, "58.3 %"] in mot        # 7/12
    assert mot[-1][1] == 12                      # TOTAL = total des demandes


def test_preaccueil_ne_propose_pas_les_services_et_motifs_inactifs(qtbot):
    from widgets.preaccueil import PreAccueilPage
    db.execute_insert("INSERT INTO accueil_service (nom, actif) VALUES ('SvcInactif', 0)")
    db.execute_insert("INSERT INTO accueil_motif (nom, actif) VALUES ('MotifInactif', 0)")
    page = PreAccueilPage()
    qtbot.addWidget(page)
    services = [page.combo_service.itemText(i) for i in range(page.combo_service.count())]
    motifs = [page.combo_motif.itemText(i) for i in range(page.combo_motif.count())]
    assert "SvcInactif" not in services and "MotifInactif" not in motifs
    assert len(services) > 1 and len(motifs) > 1  # les actifs restent proposés


def test_test_connexion_refuse_un_chemin_inexistant(tmp_path):
    from widgets.users.Main import verifier_chemin_base
    absent = tmp_path / "nouveau.sqlite3"
    ok, _ = verifier_chemin_base(str(absent))
    assert not ok and not absent.exists()        # et ne crée aucun fichier
    vide = tmp_path / "vide.sqlite3"
    import sqlite3
    sqlite3.connect(vide).close()
    assert not verifier_chemin_base(str(vide))[0]  # pas une base de l'appli
    ok, _ = verifier_chemin_base(str(db.db_path))
    assert ok


def test_id_sequentiel_survit_a_un_id_manuel_de_20_chiffres():
    """Régression: un id horodaté de 20 chiffres faisait déborder CAST(id AS INTEGER)."""
    from database.models import DemandeAccueil
    db.execute_update("DELETE FROM accueil_demandeaccueil")
    db.execute_insert("INSERT INTO accueil_demandeaccueil (id, nom, date) VALUES ('20261023121617686172', 'x', '2026-01-01')")
    ids = []
    for n in "ABC":
        d = DemandeAccueil(nom=n)
        assert d.save()
        ids.append(d.id)
    assert ids == ["944", "945", "946"]


def test_valider_demande_refuse_les_saisies_invalides():
    from widgets.database.database_page import valider_demande
    ok = {"date": "2026-05-01", "nom": "KONE", "satisfaction": None, "priorite": "haute"}
    assert valider_demande(ok) is None
    assert valider_demande({**ok, "nom": None})                       # nom obligatoire
    assert valider_demande({**ok, "date": "01/05/2026"})              # date ISO
    assert valider_demande({**ok, "date": None})
    assert valider_demande({**ok, "satisfaction": "9"})               # 1 à 5
    assert valider_demande({**ok, "satisfaction": "abc"})
    assert valider_demande({**ok, "satisfaction": "4"}) is None
    assert valider_demande({**ok, "priorite": "urgente"})


def test_seed_data_refuse_d_ecraser_une_base_existante(tmp_path):
    import sqlite3, subprocess, sys, pathlib
    base = tmp_path / "db.sqlite3"
    with sqlite3.connect(base) as c:
        c.execute("CREATE TABLE accueil_demandeaccueil (id TEXT)")
        c.execute("INSERT INTO accueil_demandeaccueil VALUES ('1')")
    script = pathlib.Path(__file__).resolve().parent.parent / "seed_data.py"
    r = subprocess.run([sys.executable, "-I", str(script)], cwd=tmp_path, capture_output=True, text=True)
    assert r.returncode != 0 and "--force" in (r.stdout + r.stderr)
    with sqlite3.connect(base) as c:
        assert c.execute("SELECT COUNT(*) FROM accueil_demandeaccueil").fetchone()[0] == 1


def test_formulaire_utilisateur_valide_identifiant_et_email(qtbot):
    import pytest
    from widgets.users.Main import UserFormDialog
    d = UserFormDialog(None, with_password=True)
    qtbot.addWidget(d)
    d.password_edit.setText("Valide#Mdp2024")
    for nom in ("a", "avec espace", "x" * 21):
        d.username_edit.setText(nom)
        with pytest.raises(ValueError):
            d.get_data()
    d.username_edit.setText("agent_01")
    d.email_edit.setText("pas-un-email")
    with pytest.raises(ValueError):
        d.get_data()
    d.email_edit.setText("")                       # e-mail optionnel
    assert d.get_data()[0] == "agent_01"
    d.email_edit.setText("agent@cnps.ci")
    assert d.get_data()[1] == "agent@cnps.ci"


def test_graduation_entiere_pas_de_doublons():
    from widgets.dashboard.Main import _graduation_entiere
    for max_val in (0, 1, 3, 5, 7, 12, 480, 5200):
        haut, pas = _graduation_entiere(max_val, 1.25)
        assert isinstance(pas, int) and pas >= 1
        assert haut % pas == 0 and haut >= max_val          # la barre la plus haute tient dans l'axe
        assert haut // pas <= 6                              # au plus 7 graduations
    assert _graduation_entiere(3, 1.25) == (4, 1)            # 0 1 2 3 4, plus de doublon
